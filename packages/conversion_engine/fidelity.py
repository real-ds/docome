import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Set, Tuple, Union

import pymupdf
from docx import Document

from .reconstruction import FontMapper

WORD_PATTERN = re.compile(r"[a-z0-9]+")


@dataclass
class FidelityReport:
    source_pages: int
    output_sections: int
    text_retention: float
    font_retention: float
    tables_source: int
    tables_output: int
    images_source: int
    images_output: int
    hyperlinks_source: int
    hyperlinks_output: int
    missing_text: List[str] = field(default_factory=list)
    missing_fonts: List[str] = field(default_factory=list)

    @property
    def page_count_preserved(self) -> bool:
        return self.source_pages == self.output_sections

    @property
    def tables_preserved(self) -> bool:
        return self.tables_output >= self.tables_source

    @property
    def images_preserved(self) -> bool:
        return self.images_output >= self.images_source

    @property
    def hyperlinks_preserved(self) -> bool:
        return self.hyperlinks_output >= self.hyperlinks_source

    @property
    def score(self) -> float:
        components = [
            self.text_retention,
            self.font_retention,
            self._ratio(self.tables_output, self.tables_source),
            self._ratio(self.images_output, self.images_source),
            self._ratio(self.hyperlinks_output, self.hyperlinks_source),
        ]
        return round(sum(components) / len(components), 4)

    @property
    def passed(self) -> bool:
        return (
            self.page_count_preserved
            and self.text_retention >= 0.9
            and self.font_retention >= 0.9
            and self.tables_preserved
            and self.images_preserved
            and self.hyperlinks_preserved
        )

    def as_dict(self) -> Dict[str, object]:
        return {
            "source_pages": self.source_pages,
            "output_sections": self.output_sections,
            "page_count_preserved": self.page_count_preserved,
            "text_retention": round(self.text_retention, 4),
            "font_retention": round(self.font_retention, 4),
            "tables": f"{self.tables_output}/{self.tables_source}",
            "tables_preserved": self.tables_preserved,
            "images": f"{self.images_output}/{self.images_source}",
            "images_preserved": self.images_preserved,
            "hyperlinks": f"{self.hyperlinks_output}/{self.hyperlinks_source}",
            "hyperlinks_preserved": self.hyperlinks_preserved,
            "score": self.score,
            "passed": self.passed,
            "missing_text": self.missing_text,
            "missing_fonts": self.missing_fonts,
        }

    @staticmethod
    def _ratio(output: int, source: int) -> float:
        if source == 0:
            return 1.0
        return min(1.0, output / source)


class ConversionFidelity:
    def evaluate(
        self,
        source_pdf: Union[str, Path],
        output_docx: Union[str, Path],
    ) -> FidelityReport:
        pdf_text, pdf_fonts, tables, images, hyperlinks, pages = self._scan_pdf(source_pdf)
        docx_text, docx_fonts, sections, docx_tables, docx_images, docx_links = self._scan_docx(output_docx)

        text_retention, missing_text = self._retention(pdf_text, docx_text)
        font_retention, missing_fonts = self._retention(pdf_fonts, docx_fonts)

        return FidelityReport(
            source_pages=pages,
            output_sections=sections,
            text_retention=text_retention,
            font_retention=font_retention,
            tables_source=tables,
            tables_output=docx_tables,
            images_source=images,
            images_output=docx_images,
            hyperlinks_source=hyperlinks,
            hyperlinks_output=docx_links,
            missing_text=missing_text,
            missing_fonts=missing_fonts,
        )

    def _scan_pdf(self, source: Union[str, Path]):
        document = pymupdf.open(str(source))
        try:
            words: Set[str] = set()
            fonts: Set[str] = set()
            tables = 0
            images: Set[int] = set()
            hyperlinks = 0

            for page in document:
                words |= tokenize(page.get_text())
                fonts |= self._pdf_fonts(page)
                try:
                    tables += len(page.find_tables().tables)
                except Exception:
                    pass
                for entry in page.get_images(full=True):
                    images.add(entry[0])
                hyperlinks += sum(
                    1 for link in page.get_links()
                    if link.get("kind") == pymupdf.LINK_URI and link.get("uri")
                )

            return words, fonts, tables, len(images), hyperlinks, document.page_count
        finally:
            document.close()

    def _pdf_fonts(self, page: pymupdf.Page) -> Set[str]:
        fonts: Set[str] = set()
        raw = page.get_text("dict", flags=pymupdf.TEXTFLAGS_DICT)
        for block in raw.get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    if not span.get("text", "").strip():
                        continue
                    fonts.add(font_key(
                        FontMapper.map_font(span.get("font", "")),
                        span.get("size", 0.0),
                    ))
        return fonts

    def _scan_docx(self, output: Union[str, Path]):
        document = Document(str(output))

        words: Set[str] = set()
        fonts: Set[str] = set()
        tables = len(document.tables)
        images = len(document.inline_shapes)
        links = sum(
            1 for rel in document.part.rels.values()
            if rel.reltype.endswith("/hyperlink")
        )

        for paragraph in document.paragraphs:
            words |= tokenize(paragraph.text)
            fonts |= self._run_fonts(paragraph.runs)

        for table in document.tables:
            words, fonts = self._scan_table(table, words, fonts)

        for section in document.sections:
            for container in (section.header, section.footer):
                for paragraph in container.paragraphs:
                    words |= tokenize(paragraph.text)
                    fonts |= self._run_fonts(paragraph.runs)

        return words, fonts, len(document.sections), tables, images, links

    def _scan_table(self, table, words: Set[str], fonts: Set[str]):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    words |= tokenize(paragraph.text)
                    fonts |= self._run_fonts(paragraph.runs)
                for nested in cell.tables:
                    words, fonts = self._scan_table(nested, words, fonts)
        return words, fonts

    def _run_fonts(self, runs: Sequence) -> Set[str]:
        fonts: Set[str] = set()
        for run in runs:
            if not run.text.strip():
                continue
            size = run.font.size.pt if run.font.size is not None else 0.0
            fonts.add(font_key(run.font.name or "", size))
        return fonts

    def _retention(self, source: Set[str], output: Set[str]) -> Tuple[float, List[str]]:
        if not source:
            return 1.0, []
        missing = sorted(source - output)
        return (len(source) - len(missing)) / len(source), missing


def tokenize(text: str) -> Set[str]:
    return set(WORD_PATTERN.findall(text.lower()))


def font_key(name: str, size: float) -> str:
    return f"{name.strip().lower()}|{round(float(size), 1)}"


def evaluate(
    source_pdf: Union[str, Path],
    output_docx: Union[str, Path],
) -> FidelityReport:
    return ConversionFidelity().evaluate(source_pdf, output_docx)
