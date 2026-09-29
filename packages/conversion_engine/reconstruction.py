import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Sequence, Tuple, Union

import pymupdf
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BOLD_FLAG = 16
ITALIC_FLAG = 2
PT_PER_CM = 28.3465


@dataclass
class FontInfo:
    name: str
    size: float
    bold: bool
    italic: bool
    color: Tuple[int, int, int]
    font_id: str = ""
    line_height: float = 0.0

    def signature(self) -> tuple:
        return (self.name, round(self.size, 1), self.bold, self.italic, self.color)


@dataclass
class TextBlock:
    x0: float
    y0: float
    x1: float
    y1: float
    text: str
    font: FontInfo
    page_num: int
    column: int = 0
    is_header: bool = False
    is_footer: bool = False
    is_list_item: bool = False
    list_style: str = ""
    list_level: int = 0
    list_marker: str = ""
    link_uri: str = ""

    @property
    def width(self) -> float:
        return self.x1 - self.x0

    @property
    def height(self) -> float:
        return self.y1 - self.y0

    @property
    def center_x(self) -> float:
        return (self.x0 + self.x1) / 2


@dataclass
class TableCell:
    x0: float
    y0: float
    x1: float
    y1: float
    text: str


@dataclass
class Table:
    x0: float
    y0: float
    x1: float
    y1: float
    rows: int
    cols: int
    cells: List[List[TableCell]] = field(default_factory=list)


@dataclass
class ImageBlock:
    x0: float
    y0: float
    x1: float
    y1: float
    image_data: bytes
    ext: str = "png"


@dataclass
class LinkBlock:
    x0: float
    y0: float
    x1: float
    y1: float
    uri: str
    text: str = ""


@dataclass
class ColumnRegion:
    x0: float
    x1: float
    blocks: List[TextBlock] = field(default_factory=list)


@dataclass
class PageStructure:
    page_num: int
    width: float
    height: float
    text_blocks: List[TextBlock] = field(default_factory=list)
    tables: List[Table] = field(default_factory=list)
    images: List[ImageBlock] = field(default_factory=list)
    links: List[LinkBlock] = field(default_factory=list)
    columns: List[ColumnRegion] = field(default_factory=list)
    margin_left: float = 72.0
    margin_right: float = 72.0
    margin_top: float = 72.0
    margin_bottom: float = 72.0

    @property
    def column_count(self) -> int:
        return max(1, len(self.columns))

    @property
    def content_width(self) -> float:
        return self.width - self.margin_left - self.margin_right

    def in_reading_order(self) -> List[TextBlock]:
        return sorted(
            self.text_blocks,
            key=lambda b: (b.column, round(b.y0, 1), b.x0),
        )


class FontMapper:
    FONT_ALIASES = (
        ("times", "Times New Roman"),
        ("roman", "Times New Roman"),
        ("georgia", "Times New Roman"),
        ("courier", "Courier New"),
        ("mono", "Courier New"),
        ("symbol", "Symbol"),
        ("zapf", "Wingdings"),
        ("arial", "Arial"),
        ("helv", "Arial"),
    )
    VARIANT_SUFFIXES = ("bold", "italic", "oblique", "boldoblique", "bolditalic")

    @classmethod
    def is_bold(cls, flags: int, pdf_font_name: str) -> bool:
        if flags & BOLD_FLAG:
            return True
        lowered = pdf_font_name.lower()
        return any("bold" in variant for variant in cls.VARIANT_SUFFIXES if variant in lowered)

    @classmethod
    def is_italic(cls, flags: int, pdf_font_name: str) -> bool:
        if flags & ITALIC_FLAG:
            return True
        lowered = pdf_font_name.lower()
        return "italic" in lowered or "oblique" in lowered

    @classmethod
    def map_font(cls, pdf_font_name: str) -> str:
        cleaned = re.sub(r"[-+,](Bold|Italic|Oblique|BoldOblique|BoldItalic)$", "", pdf_font_name, flags=re.IGNORECASE)
        lowered = cleaned.lower().replace(" ", "").replace("-", "")
        for alias, mapped in cls.FONT_ALIASES:
            if alias in lowered:
                return mapped
        return cleaned.strip() or "Arial"

    @staticmethod
    def color_to_rgb(color: int) -> Tuple[int, int, int]:
        value = int(color) & 0xFFFFFF
        return ((value >> 16) & 0xFF, (value >> 8) & 0xFF, value & 0xFF)


class ListDetector:
    BULLET_PATTERN = re.compile(r"^([\u2022\u2023\u2043\u2219\u25cf\u25aa\u25e6\u25c6\u25b8\u00b7*+\-\u2013\u2014])\s+(.*)$")
    NUMBER_PATTERN = re.compile(r"^(\(?\d{1,3}[.)]|\(?[a-zA-Z][.)]|\(?[ivxlcdmIVXLCDM]{1,6}[.)])\s+(.*)$")
    INDENT_STEP = 18.0

    def annotate(self, block: TextBlock, left_margin: float) -> TextBlock:
        bullet = self.BULLET_PATTERN.match(block.text)
        if bullet:
            block.is_list_item = True
            block.list_style = "bullet"
            block.list_marker = bullet.group(1)
            block.text = bullet.group(2).strip()
            block.list_level = self._level(block, left_margin)
            return block

        numbered = self.NUMBER_PATTERN.match(block.text)
        if numbered:
            block.is_list_item = True
            block.list_style = "number"
            block.list_marker = numbered.group(1)
            block.text = numbered.group(2).strip()
            block.list_level = self._level(block, left_margin)
        return block

    def _level(self, block: TextBlock, left_margin: float) -> int:
        offset = max(0.0, block.x0 - left_margin)
        return min(8, int(offset // self.INDENT_STEP))


class TextBlockDetector:
    def __init__(self, line_height_factor: float = 1.9, left_edge_tolerance: float = 4.0, right_fill_ratio: float = 0.82):
        self.line_height_factor = line_height_factor
        self.left_edge_tolerance = left_edge_tolerance
        self.right_fill_ratio = right_fill_ratio

    def detect_lines(self, page: pymupdf.Page, page_num: int) -> List[TextBlock]:
        raw = page.get_text("dict", flags=pymupdf.TEXTFLAGS_DICT)
        lines: List[TextBlock] = []

        for block in raw.get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                spans = line.get("spans", [])
                if not spans:
                    continue

                text = "".join(span.get("text", "") for span in spans)
                if not text.strip():
                    continue

                font = self._font_from_spans(spans)
                bbox = line.get("bbox") or (0, 0, 0, 0)
                lines.append(TextBlock(
                    x0=bbox[0], y0=bbox[1], x1=bbox[2], y1=bbox[3],
                    text=text.strip(),
                    font=font,
                    page_num=page_num,
                ))

        lines.sort(key=lambda b: (round(b.y0, 1), b.x0))
        return lines

    def merge_paragraphs(self, lines: Sequence[TextBlock], column_right: float, left_margin: float) -> List[TextBlock]:
        if not lines:
            return []

        list_detector = ListDetector()
        annotated = [list_detector.annotate(line, left_margin) for line in lines]

        merged: List[TextBlock] = []
        current = annotated[0]

        for candidate in annotated[1:]:
            if self._joins(current, candidate, column_right):
                current.text = f"{current.text} {candidate.text}".strip()
                current.x1 = max(current.x1, candidate.x1)
                current.y1 = max(current.y1, candidate.y1)
            else:
                merged.append(current)
                current = candidate

        merged.append(current)
        return merged

    def _joins(self, current: TextBlock, candidate: TextBlock, column_right: float) -> bool:
        if current.column != candidate.column:
            return False
        if current.is_list_item or candidate.is_list_item:
            return False
        if current.font.signature() != candidate.font.signature():
            return False
        if abs(candidate.x0 - current.x0) > self.left_edge_tolerance:
            return False

        gap = candidate.y0 - current.y1
        if gap < -1.0 or gap > self.line_height_factor * max(current.font.size, 1.0):
            return False

        if current.x1 < column_right * self.right_fill_ratio:
            return False
        return True

    def _font_from_spans(self, spans: Sequence[dict]) -> FontInfo:
        primary = max(spans, key=lambda span: len(span.get("text", "").strip()))
        pdf_font_name = primary.get("font", "Helvetica")
        size = float(primary.get("size", 12.0))
        return FontInfo(
            name=FontMapper.map_font(pdf_font_name),
            size=size,
            bold=FontMapper.is_bold(primary.get("flags", 0), pdf_font_name),
            italic=FontMapper.is_italic(primary.get("flags", 0), pdf_font_name),
            color=FontMapper.color_to_rgb(primary.get("color", 0)),
            font_id=pdf_font_name,
            line_height=size * self.line_height_factor,
        )


class ColumnDetector:
    def __init__(self, gap_ratio: float = 0.06, min_blocks_per_column: int = 2):
        self.gap_ratio = gap_ratio
        self.min_blocks_per_column = min_blocks_per_column

    def detect(self, blocks: Sequence[TextBlock], page_width: float) -> List[ColumnRegion]:
        for block in blocks:
            block.column = 0

        if not blocks:
            return []

        groups = self._cluster_by_x(blocks, page_width * self.gap_ratio)

        if len(groups) < 2 or any(len(group) < self.min_blocks_per_column for group in groups):
            return [self._region(blocks)]

        regions = [self._region(group) for group in groups]
        for index, region in enumerate(regions):
            for block in region.blocks:
                block.column = index
        return regions

    def _cluster_by_x(self, blocks: Sequence[TextBlock], gap: float) -> List[List[TextBlock]]:
        ordered = sorted(blocks, key=lambda block: (block.x0, block.y0))
        groups: List[List[TextBlock]] = []
        current: List[TextBlock] = []
        current_right = 0.0

        for block in ordered:
            if current and block.x0 > current_right + gap:
                groups.append(current)
                current = []
                current_right = 0.0
            current.append(block)
            current_right = max(current_right, block.x1)

        if current:
            groups.append(current)
        return groups

    def _region(self, blocks: Sequence[TextBlock]) -> ColumnRegion:
        return ColumnRegion(
            x0=min(block.x0 for block in blocks),
            x1=max(block.x1 for block in blocks),
            blocks=list(blocks),
        )


class HeaderFooterDetector:
    PAGE_NUMBER_PATTERN = re.compile(r"^\s*(page\s+)?\d+\s*(of|/)\s*\d+\s*$", re.IGNORECASE)
    BARE_NUMBER_PATTERN = re.compile(r"^\s*-?\s*\d{1,4}\s*-?\s*$")
    SIZE_MARGIN = 0.5

    def __init__(self, header_ratio: float = 0.09, footer_ratio: float = 0.91):
        self.header_ratio = header_ratio
        self.footer_ratio = footer_ratio

    def annotate(self, blocks: Sequence[TextBlock], page_height: float) -> None:
        header_limit = page_height * self.header_ratio
        footer_limit = page_height * self.footer_ratio
        body_size = self._body_font_size(blocks, header_limit, footer_limit)
        if body_size is None:
            return

        for block in blocks:
            in_header_band = block.y1 <= header_limit
            in_footer_band = block.y0 >= footer_limit

            if not (in_header_band or in_footer_band):
                continue

            if self._looks_like_page_number(block.text) or block.font.size <= body_size - self.SIZE_MARGIN:
                block.is_header = in_header_band
                block.is_footer = in_footer_band

    def _body_font_size(self, blocks: Sequence[TextBlock], header_limit: float, footer_limit: float) -> Optional[float]:
        sizes = [
            block.font.size
            for block in blocks
            if block.y1 > header_limit and block.y0 < footer_limit
        ]
        return max(sizes) if sizes else None

    def _looks_like_page_number(self, text: str) -> bool:
        return bool(self.PAGE_NUMBER_PATTERN.match(text) or self.BARE_NUMBER_PATTERN.match(text))


class TableDetector:
    def __init__(self, min_rows: int = 2, min_cols: int = 2):
        self.min_rows = min_rows
        self.min_cols = min_cols

    def detect(self, page: pymupdf.Page) -> List[Table]:
        try:
            found = page.find_tables()
        except Exception:
            return []

        tables: List[Table] = []
        for item in found.tables:
            data = item.extract()
            if len(data) < self.min_rows:
                continue

            columns = max((len(row) for row in data), default=0)
            if columns < self.min_cols:
                continue

            bbox = item.bbox
            cells = self._build_cells(data, item)
            tables.append(Table(
                x0=bbox[0], y0=bbox[1], x1=bbox[2], y1=bbox[3],
                rows=len(data), cols=columns, cells=cells,
            ))
        return tables

    def _build_cells(self, data: Sequence[Sequence[Optional[str]]], item) -> List[List[TableCell]]:
        cells: List[List[TableCell]] = []
        for row_index, row in enumerate(data):
            rects = self._row_rects(item, row_index)
            built_row: List[TableCell] = []
            for col_index, value in enumerate(row):
                rect = rects[col_index] if col_index < len(rects) and rects[col_index] else (0.0, 0.0, 0.0, 0.0)
                built_row.append(TableCell(
                    x0=rect[0], y0=rect[1], x1=rect[2], y1=rect[3],
                    text=(value or "").strip(),
                ))
            cells.append(built_row)
        return cells

    def _row_rects(self, item, row_index: int) -> List[Tuple[float, float, float, float]]:
        try:
            return [tuple(rect) for rect in item.rows[row_index].cells if rect is not None]
        except (AttributeError, IndexError):
            return []


class ImageExtractor:
    def extract(self, page: pymupdf.Page) -> List[ImageBlock]:
        images: List[ImageBlock] = []
        seen: set = set()

        for entry in page.get_images(full=True):
            xref = entry[0]
            if xref in seen:
                continue
            seen.add(xref)

            for rect in self._rects(page, xref):
                payload = self._payload(page, xref)
                if payload is None:
                    continue
                data, extension = payload
                images.append(ImageBlock(
                    x0=rect[0], y0=rect[1], x1=rect[2], y1=rect[3],
                    image_data=data, ext=extension,
                ))
        images.sort(key=lambda image: (image.y0, image.x0))
        return images

    def _rects(self, page: pymupdf.Page, xref: int) -> List[Tuple[float, float, float, float]]:
        try:
            return [tuple(rect) for rect in page.get_image_rects(xref)]
        except Exception:
            return []

    def _payload(self, page: pymupdf.Page, xref: int):
        try:
            extracted = page.parent.extract_image(xref)
        except Exception:
            return None
        if not extracted:
            return None
        return extracted.get("image"), extracted.get("ext", "png")


class LinkExtractor:
    def extract(self, page: pymupdf.Page) -> List[LinkBlock]:
        links: List[LinkBlock] = []
        for entry in page.get_links():
            if entry.get("kind") != pymupdf.LINK_URI:
                continue
            uri = entry.get("uri")
            if not uri:
                continue
            rect = entry.get("from")
            if rect is None:
                continue
            links.append(LinkBlock(
                x0=rect[0], y0=rect[1], x1=rect[2], y1=rect[3],
                uri=uri,
                text=self._text_for(page, rect),
            ))
        return links

    def _text_for(self, page: pymupdf.Page, rect) -> str:
        try:
            return page.get_textbox(rect).strip()
        except Exception:
            return ""

    @staticmethod
    def attach(blocks: Sequence[TextBlock], links: Sequence[LinkBlock]) -> None:
        for block in blocks:
            best: Optional[LinkBlock] = None
            best_score = 0.0
            for link in links:
                score = LinkExtractor._overlap(block, link)
                if score > best_score:
                    best_score = score
                    best = link
            if best is not None and best_score >= 0.6:
                block.link_uri = best.uri

    @staticmethod
    def _overlap(block: TextBlock, link: LinkBlock) -> float:
        x_overlap = min(block.x1, link.x1) - max(block.x0, link.x0)
        y_overlap = min(block.y1, link.y1) - max(block.y0, link.y0)
        if x_overlap <= 0 or y_overlap <= 0:
            return 0.0
        link_area = max((link.x1 - link.x0) * (link.y1 - link.y0), 1e-6)
        return (x_overlap * y_overlap) / link_area


class PageLayout:
    def __init__(self, min_margin: float = 36.0, max_margin: float = 90.0):
        self.min_margin = min_margin
        self.max_margin = max_margin

    def margins(self, structure: PageStructure) -> None:
        body = [block for block in structure.text_blocks if not block.is_header and not block.is_footer]
        sources = body or structure.text_blocks

        if sources:
            left = min(block.x0 for block in sources)
            right = structure.width - max(block.x1 for block in sources)
            top = min(block.y0 for block in sources)
            bottom = structure.height - max(block.y1 for block in sources)
        else:
            left = right = top = bottom = self.max_margin

        structure.margin_left = self._clamp(left)
        structure.margin_right = self._clamp(right)
        structure.margin_top = self._clamp(top)
        structure.margin_bottom = self._clamp(bottom)

    def _clamp(self, value: float) -> float:
        return max(self.min_margin, min(self.max_margin, value))


class DocxBuilder:
    LIST_STYLES = {"bullet": "List Bullet", "number": "List Number"}

    def __init__(self, document: Document):
        self.document = document

    def start_page(self, structure: PageStructure, first: bool) -> None:
        section = self.document.sections[0] if first else self.document.add_section(WD_SECTION_START.NEW_PAGE)
        self._apply_page_setup(section, structure)
        self._apply_columns(section, structure.column_count)

    def add_paragraph(self, block: TextBlock, structure: PageStructure) -> None:
        if block.is_header or block.is_footer:
            return

        style = self.LIST_STYLES.get(block.list_style) if block.is_list_item else None
        paragraph = self.document.add_paragraph(style=style) if style else self.document.add_paragraph()

        text = f"{block.list_marker} {block.text}".strip() if block.is_list_item and style is None else block.text
        if block.link_uri:
            self._add_hyperlink(paragraph, text, block.link_uri, block.font)
        else:
            self._add_run(paragraph, text, block.font)

        alignment = self._alignment(block, structure)
        if alignment is not None:
            paragraph.alignment = alignment

        if block.is_list_item and block.list_level:
            self._indent(paragraph, block.list_level * 0.5)

    def add_table(self, table: Table) -> None:
        if not table.cells or table.rows < 1 or table.cols < 1:
            return
        docx_table = self.document.add_table(rows=table.rows, cols=table.cols)
        docx_table.style = "Table Grid"
        for row_index, row in enumerate(table.cells):
            for col_index, cell in enumerate(row):
                if col_index >= len(docx_table.rows[row_index].cells):
                    continue
                target = docx_table.rows[row_index].cells[col_index]
                target.text = cell.text
                for paragraph in target.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        run.font.size = Pt(10)

    def add_image(self, image: ImageBlock) -> None:
        width_pt = image.x1 - image.x0
        if width_pt <= 0:
            return
        try:
            self.document.add_picture(
                io.BytesIO(image.image_data),
                width=Cm(width_pt / PT_PER_CM),
            )
        except Exception:
            return

    def set_header_footer(self, section, headers: Sequence[str], footers: Sequence[str]) -> None:
        self._write_side(section.header, headers)
        self._write_side(section.footer, footers)

    def _write_side(self, container, texts: Sequence[str]) -> None:
        if not texts:
            return
        container.is_linked_to_previous = False
        paragraph = container.paragraphs[0] if container.paragraphs else container.add_paragraph()
        for run in list(paragraph.runs):
            run._r.getparent().remove(run._r)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._add_run(paragraph, " | ".join(texts), FontInfo("Arial", 9.0, False, False, (0, 0, 0)))

    def _add_run(self, paragraph, text: str, font: FontInfo):
        run = paragraph.add_run(text)
        run.font.name = font.name
        run.font.size = Pt(font.size)
        run.font.bold = font.bold
        run.font.italic = font.italic
        run.font.color.rgb = RGBColor(*font.color)
        return run

    def _add_hyperlink(self, paragraph, text: str, uri: str, font: FontInfo) -> None:
        run = self._add_run(paragraph, text, font)
        relationship_id = paragraph.part.relate_to(uri, RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship_id)
        paragraph._p.remove(run._r)
        hyperlink.append(run._r)
        paragraph._p.append(hyperlink)

    def _alignment(self, block: TextBlock, structure: PageStructure) -> Optional[WD_ALIGN_PARAGRAPH]:
        tolerance = structure.content_width * 0.08
        left_gap = abs(block.x0 - structure.margin_left)
        right_gap = abs((structure.width - block.x1) - structure.margin_right)

        if abs(block.center_x - structure.width / 2) <= tolerance:
            return WD_ALIGN_PARAGRAPH.CENTER
        if right_gap <= tolerance and right_gap < left_gap:
            return WD_ALIGN_PARAGRAPH.RIGHT
        if left_gap <= max(tolerance, 12.0):
            return WD_ALIGN_PARAGRAPH.LEFT
        return None

    def _indent(self, paragraph, level: float) -> None:
        paragraph.paragraph_format.left_indent = Cm(level)

    def _apply_page_setup(self, section, structure: PageStructure) -> None:
        section.page_width = Pt(structure.width)
        section.page_height = Pt(structure.height)
        section.left_margin = Pt(structure.margin_left)
        section.right_margin = Pt(structure.margin_right)
        section.top_margin = Pt(structure.margin_top)
        section.bottom_margin = Pt(structure.margin_bottom)

    def _apply_columns(self, section, count: int) -> None:
        columns = section._sectPr.find(qn("w:cols"))
        if columns is None:
            columns = OxmlElement("w:cols")
            section._sectPr.append(columns)
        columns.set(qn("w:num"), str(max(1, count)))
        columns.set(qn("w:equalWidth"), "1")
        columns.set(qn("w:space"), str(int(0.5 * 72 * 20)))


class HighFidelityPdfToDocx:
    def __init__(self):
        self.text_detector = TextBlockDetector()
        self.column_detector = ColumnDetector()
        self.header_footer_detector = HeaderFooterDetector()
        self.table_detector = TableDetector()
        self.image_extractor = ImageExtractor()
        self.link_extractor = LinkExtractor()
        self.layout = PageLayout()

    def convert(
        self,
        input_path: Union[str, Path, bytes],
        output_path: Union[str, Path],
    ) -> Document:
        pdf_document = self._open(input_path)
        try:
            document = self.build(pdf_document)
            document.save(str(output_path))
            return document
        finally:
            pdf_document.close()

    def build(self, pdf_document: pymupdf.Document) -> Document:
        document = Document()
        self._reset_body(document)
        builder = DocxBuilder(document)

        for page_number in range(pdf_document.page_count):
            structure = self.analyze_page(pdf_document[page_number], page_number)
            builder.start_page(structure, first=page_number == 0)
            builder.set_header_footer(
                document.sections[-1],
                [block.text for block in structure.text_blocks if block.is_header],
                [block.text for block in structure.text_blocks if block.is_footer],
            )
            for block in structure.in_reading_order():
                builder.add_paragraph(block, structure)
            for table in structure.tables:
                builder.add_table(table)
            for image in structure.images:
                builder.add_image(image)

        return document

    def analyze_page(self, page: pymupdf.Page, page_number: int) -> PageStructure:
        structure = PageStructure(
            page_num=page_number,
            width=page.rect.width,
            height=page.rect.height,
            tables=self.table_detector.detect(page),
            images=self.image_extractor.extract(page),
            links=self.link_extractor.extract(page),
        )

        lines = self.text_detector.detect_lines(page, page_number)
        body = [line for line in lines if not self._in_table(line, structure.tables)]
        prelim_columns = self.column_detector.detect(body, structure.width)

        left_margin = min((block.x0 for block in body), default=0.0)
        structure.text_blocks = self.text_detector.merge_paragraphs(
            body,
            column_right=prelim_columns[-1].x1 if prelim_columns else structure.width,
            left_margin=left_margin,
        )
        self.header_footer_detector.annotate(structure.text_blocks, structure.height)
        self._mark_header_footer_columns(structure)
        LinkExtractor.attach(structure.text_blocks, structure.links)

        self.layout.margins(structure)
        structure.columns = self.column_detector.detect(
            self._column_source(structure), structure.width
        )
        return structure

    def _in_table(self, line: TextBlock, tables: Sequence[Table]) -> bool:
        for table in tables:
            if line.x0 >= table.x0 - 2 and line.x1 <= table.x1 + 2 and line.y0 >= table.y0 - 2 and line.y1 <= table.y1 + 2:
                return True
        return False

    def _mark_header_footer_columns(self, structure: PageStructure) -> None:
        for block in structure.text_blocks:
            if block.is_header or block.is_footer:
                block.column = 0

    def _column_source(self, structure: PageStructure) -> List[TextBlock]:
        return [
            block for block in structure.text_blocks
            if not block.is_header and not block.is_footer
        ]

    def _open(self, input_path: Union[str, Path, bytes]) -> pymupdf.Document:
        if isinstance(input_path, (bytes, bytearray, memoryview)):
            return pymupdf.open(stream=bytes(input_path), filetype="pdf")
        return pymupdf.open(str(input_path))

    @staticmethod
    def _reset_body(document: Document) -> None:
        body = document.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)


def pdf_to_docx_hq(
    input_path: Union[str, Path, bytes],
    output_path: Union[str, Path],
) -> Document:
    return HighFidelityPdfToDocx().convert(input_path, output_path)
