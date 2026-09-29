import io

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn

from packages.conversion_engine.reconstruction import (
    ColumnDetector,
    DocxBuilder,
    FontInfo,
    FontMapper,
    HeaderFooterDetector,
    HighFidelityPdfToDocx,
    ImageExtractor,
    LinkExtractor,
    ListDetector,
    PageStructure,
    Table,
    TableDetector,
    TextBlock,
    TextBlockDetector,
    pdf_to_docx_hq,
)

PAGE_WIDTH = 595.0
PAGE_HEIGHT = 842.0


def make_block(text, x0=72.0, y0=72.0, x1=300.0, y1=86.0, font=None, page_num=0):
    return TextBlock(
        x0=x0, y0=y0, x1=x1, y1=y1, text=text,
        font=font or FontInfo("Arial", 12.0, False, False, (0, 0, 0)),
        page_num=page_num,
    )


@pytest.fixture
def simple_pdf_path(tmp_path):
    document = pymupdf.open()
    page = document.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    page.insert_text((72, 120), "Chapter One", fontname="hebo", fontsize=20)
    page.insert_text((72, 160), "Body text that continues", fontname="helv", fontsize=11)
    page.insert_text((72, 176), "onto a second visual line.", fontname="helv", fontsize=11)
    path = tmp_path / "simple.pdf"
    document.save(str(path))
    document.close()
    return path


@pytest.fixture
def multi_page_pdf_path(tmp_path):
    document = pymupdf.open()
    for index in range(3):
        page = document.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        page.insert_text((72, 50), f"Running head {index + 1}", fontname="helv", fontsize=8)
        page.insert_text((72, 140), f"Page {index + 1} content", fontname="helv", fontsize=12)
        page.insert_text((300, 815), "12", fontname="helv", fontsize=8)
    path = tmp_path / "multi.pdf"
    document.save(str(path))
    document.close()
    return path


@pytest.fixture
def rich_pdf_path(tmp_path):
    document = pymupdf.open()
    page = document.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
    page.insert_text((72, 110), "Quarterly Summary", fontname="hebo", fontsize=18)
    page.insert_text((72, 150), "\u2022 First bullet item", fontname="helv", fontsize=11)
    page.insert_text((90, 168), "\u2022 Nested bullet item", fontname="helv", fontsize=11)
    page.insert_text((72, 190), "1. Numbered item", fontname="helv", fontsize=11)
    page.insert_text((72, 250), "Left column text", fontname="helv", fontsize=10)
    page.insert_text((320, 250), "Right column text", fontname="helv", fontsize=10)
    page.insert_text((72, 266), "Left column tail", fontname="helv", fontsize=10)
    page.insert_text((320, 266), "Right column tail", fontname="helv", fontsize=10)

    shapes = page.draw_rect(pymupdf.Rect(72, 320, 400, 400))
    page.draw_line(pymupdf.Point(72, 360), pymupdf.Point(400, 360))
    page.draw_line(pymupdf.Point(236, 320), pymupdf.Point(236, 400))
    page.insert_text((80, 345), "Cell A1", fontname="helv", fontsize=9)
    page.insert_text((250, 345), "Cell B1", fontname="helv", fontsize=9)
    page.insert_text((80, 385), "Cell A2", fontname="helv", fontsize=9)
    page.insert_text((250, 385), "Cell B2", fontname="helv", fontsize=9)

    page.insert_text((72, 430), "Docome project", fontname="helv", fontsize=11)
    page.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(72, 420, 150, 434), "uri": "https://example.com/docome"})

    image_bytes = _png_bytes()
    page.insert_image(pymupdf.Rect(72, 450, 200, 540), stream=image_bytes)

    path = tmp_path / "rich.pdf"
    document.save(str(path))
    document.close()
    return path


def _png_bytes():
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (24, 24), (10, 120, 200)).save(buffer, format="PNG")
    return buffer.getvalue()


class TestFontMapper:
    def test_maps_standard_families(self):
        assert FontMapper.map_font("Helvetica") == "Arial"
        assert FontMapper.map_font("Times-Roman") == "Times New Roman"
        assert FontMapper.map_font("CourierNewPSMT") == "Courier New"

    def test_strips_variant_suffixes(self):
        assert FontMapper.map_font("Helvetica-BoldOblique") == "Arial"

    def test_bold_uses_flag_16(self):
        assert FontMapper.is_bold(16, "Helvetica") is True
        assert FontMapper.is_bold(0, "Helvetica") is False

    def test_italic_uses_flag_2(self):
        assert FontMapper.is_italic(2, "Helvetica") is True
        assert FontMapper.is_italic(16, "Helvetica-Bold") is False

    def test_detects_bold_and_italic_from_font_name(self):
        assert FontMapper.is_bold(0, "Helvetica-Bold") is True
        assert FontMapper.is_italic(0, "Times-Italic") is True

    def test_color_to_rgb(self):
        assert FontMapper.color_to_rgb(0) == (0, 0, 0)
        assert FontMapper.color_to_rgb(0xFF8000) == (255, 128, 0)


class TestListDetector:
    def test_detects_bullet(self):
        block = make_block("\u2022 Bullet text")
        ListDetector().annotate(block, 72.0)
        assert block.is_list_item is True
        assert block.list_style == "bullet"
        assert block.text == "Bullet text"

    def test_detects_numbered(self):
        block = make_block("3. Third item")
        ListDetector().annotate(block, 72.0)
        assert block.list_style == "number"
        assert block.list_marker == "3."
        assert block.text == "Third item"

    def test_ignores_plain_text(self):
        block = make_block("Regular sentence.")
        ListDetector().annotate(block, 72.0)
        assert block.is_list_item is False

    def test_level_scales_with_indent(self):
        block = make_block("\u2022 Nested", x0=126.0)
        ListDetector().annotate(block, 72.0)
        assert block.list_level == 3


class TestTextBlockDetector:
    def test_merges_wrapped_lines_into_paragraph(self, simple_pdf_path):
        document = pymupdf.open(str(simple_pdf_path))
        detector = TextBlockDetector()
        lines = detector.detect_lines(document[0], 0)
        body = [line for line in lines if "Body text" in line.text or "second visual" in line.text]
        column_right = max(line.x1 for line in body)
        paragraphs = detector.merge_paragraphs(body, column_right=column_right, left_margin=72.0)
        document.close()

        assert len(paragraphs) == 1
        assert paragraphs[0].text == "Body text that continues onto a second visual line."

    def test_pipeline_joins_wrapped_lines(self, simple_pdf_path):
        converter = HighFidelityPdfToDocx()
        document = pymupdf.open(str(simple_pdf_path))
        structure = converter.analyze_page(document[0], 0)
        document.close()
        assert any(
            block.text == "Body text that continues onto a second visual line."
            for block in structure.text_blocks
        )

    def test_keeps_distinct_font_paragraphs_separate(self, simple_pdf_path):
        document = pymupdf.open(str(simple_pdf_path))
        detector = TextBlockDetector()
        lines = detector.detect_lines(document[0], 0)
        document.close()
        assert any("Chapter One" in line.text for line in lines)

    def test_reads_font_properties_from_spans(self, simple_pdf_path):
        document = pymupdf.open(str(simple_pdf_path))
        detector = TextBlockDetector()
        lines = detector.detect_lines(document[0], 0)
        document.close()

        heading = next(line for line in lines if "Chapter One" in line.text)
        assert heading.font.bold is True
        assert round(heading.font.size, 1) == 20.0

    def test_detect_lines_skips_blank_lines(self, simple_pdf_path):
        document = pymupdf.open(str(simple_pdf_path))
        lines = TextBlockDetector().detect_lines(document[0], 0)
        document.close()
        assert all(line.text.strip() for line in lines)


class TestColumnDetector:
    def test_single_column_returns_one_region(self):
        blocks = [make_block("a", y0=100), make_block("b", y0=130), make_block("c", y0=160)]
        regions = ColumnDetector().detect(blocks, PAGE_WIDTH)
        assert len(regions) == 1
        assert all(block.column == 0 for block in blocks)

    def test_two_columns_are_split(self):
        blocks = [
            make_block("l1", x0=72, x1=260, y0=100),
            make_block("l2", x0=72, x1=260, y0=130),
            make_block("r1", x0=320, x1=500, y0=100),
            make_block("r2", x0=320, x1=500, y0=130),
        ]
        regions = ColumnDetector().detect(blocks, PAGE_WIDTH)
        assert len(regions) == 2
        assert [block.column for block in blocks] == [0, 0, 1, 1]

    def test_reading_order_follows_columns(self):
        blocks = [
            make_block("r1", x0=320, x1=500, y0=100),
            make_block("l1", x0=72, x1=260, y0=100),
            make_block("r2", x0=320, x1=500, y0=130),
            make_block("l2", x0=72, x1=260, y0=130),
        ]
        for block in blocks:
            block.column = 1 if block.x0 > 300 else 0
        structure = PageStructure(page_num=0, width=PAGE_WIDTH, height=PAGE_HEIGHT, text_blocks=blocks)
        assert [block.text for block in structure.in_reading_order()] == ["l1", "l2", "r1", "r2"]


class TestHeaderFooterDetector:
    def test_marks_top_block_as_header(self):
        header = make_block("Running head", y0=40, y1=50, font=FontInfo("Arial", 8.0, False, False, (0, 0, 0)))
        body = make_block("Body content", y0=400, y1=412, font=FontInfo("Arial", 12.0, False, False, (0, 0, 0)))
        HeaderFooterDetector().annotate([header, body], PAGE_HEIGHT)
        assert header.is_header is True
        assert body.is_header is False

    def test_marks_bottom_number_as_footer(self):
        footer = make_block("12", y0=800, y1=810, font=FontInfo("Arial", 8.0, False, False, (0, 0, 0)))
        body = make_block("Body content", y0=400, y1=412, font=FontInfo("Arial", 12.0, False, False, (0, 0, 0)))
        HeaderFooterDetector().annotate([footer, body], PAGE_HEIGHT)
        assert footer.is_footer is True

    def test_body_block_is_untouched(self):
        block = make_block("Body", y0=400, y1=412)
        HeaderFooterDetector().annotate([block], PAGE_HEIGHT)
        assert block.is_header is False
        assert block.is_footer is False

    def test_large_top_block_is_not_header(self):
        heading = make_block("INVOICE #2041", y0=40, y1=60, font=FontInfo("Arial", 18.0, True, False, (0, 0, 0)))
        date = make_block("2026-09-30", x0=300, x1=360, y0=40, y1=52, font=FontInfo("Arial", 10.0, False, False, (0, 0, 0)))
        body = make_block("Line item", y0=400, y1=412, font=FontInfo("Arial", 12.0, False, False, (0, 0, 0)))
        HeaderFooterDetector().annotate([heading, date, body], PAGE_HEIGHT)
        assert heading.is_header is False
        assert date.is_header is True

    def test_bare_number_in_body_stays_body(self):
        block = make_block("12", y0=400, y1=410)
        body = make_block("Body content", y0=300, y1=312)
        HeaderFooterDetector().annotate([block, body], PAGE_HEIGHT)
        assert block.is_footer is False

    def test_page_without_body_keeps_all_content(self):
        only = make_block("Standalone title", y0=40, y1=60)
        HeaderFooterDetector().annotate([only], PAGE_HEIGHT)
        assert only.is_header is False
        assert only.is_footer is False


class TestTableDetector:
    def test_detects_drawn_table(self, rich_pdf_path):
        document = pymupdf.open(str(rich_pdf_path))
        tables = TableDetector().detect(document[0])
        document.close()

        assert len(tables) == 1
        table = tables[0]
        assert table.rows == 2
        assert table.cols == 2
        assert table.cells[0][0].text == "Cell A1"
        assert table.cells[1][1].text == "Cell B2"

    def test_returns_empty_without_tables(self, simple_pdf_path):
        document = pymupdf.open(str(simple_pdf_path))
        tables = TableDetector().detect(document[0])
        document.close()
        assert tables == []


class TestImageExtractor:
    def test_extracts_embedded_image(self, rich_pdf_path):
        document = pymupdf.open(str(rich_pdf_path))
        images = ImageExtractor().extract(document[0])
        document.close()
        assert len(images) == 1
        assert images[0].image_data[:4] == b"\x89PNG"


class TestLinkExtractor:
    def test_extracts_uri_and_text(self, rich_pdf_path):
        document = pymupdf.open(str(rich_pdf_path))
        links = LinkExtractor().extract(document[0])
        document.close()
        assert len(links) == 1
        assert links[0].uri == "https://example.com/docome"
        assert "Docome" in links[0].text

    def test_attaches_uri_to_matching_block(self, rich_pdf_path):
        document = pymupdf.open(str(rich_pdf_path))
        links = LinkExtractor().extract(document[0])
        block = make_block("Docome project", x0=72, y0=420, x1=150, y1=434)
        LinkExtractor.attach([block], links)
        document.close()
        assert block.link_uri == "https://example.com/docome"


class TestPageAnalysis:
    def test_structure_reports_two_columns(self, rich_pdf_path):
        converter = HighFidelityPdfToDocx()
        document = pymupdf.open(str(rich_pdf_path))
        structure = converter.analyze_page(document[0], 0)
        document.close()
        assert structure.column_count == 2

    def test_list_items_detected_in_page(self, rich_pdf_path):
        converter = HighFidelityPdfToDocx()
        document = pymupdf.open(str(rich_pdf_path))
        structure = converter.analyze_page(document[0], 0)
        document.close()
        styles = {block.list_style for block in structure.text_blocks if block.is_list_item}
        assert styles == {"bullet", "number"}

    def test_table_text_excluded_from_paragraphs(self, rich_pdf_path):
        converter = HighFidelityPdfToDocx()
        document = pymupdf.open(str(rich_pdf_path))
        structure = converter.analyze_page(document[0], 0)
        document.close()
        body = " ".join(block.text for block in structure.text_blocks)
        assert "Cell A1" not in body

    def test_margins_are_derived_from_content(self, simple_pdf_path):
        converter = HighFidelityPdfToDocx()
        document = pymupdf.open(str(simple_pdf_path))
        structure = converter.analyze_page(document[0], 0)
        document.close()
        assert 36.0 <= structure.margin_left <= 90.0
        assert structure.content_width > 0

    def test_header_blocks_do_not_create_columns(self, tmp_path):
        document = pymupdf.open()
        page = document.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        page.insert_text((72, 140), "Main body line one", fontname="helv", fontsize=12)
        page.insert_text((72, 160), "Main body line two", fontname="helv", fontsize=12)
        page.insert_text((72, 50), "Report Title", fontname="helv", fontsize=9)
        page.insert_text((300, 815), "Page 7", fontname="helv", fontsize=9)
        path = tmp_path / "single.pdf"
        document.save(str(path))
        document.close()

        converter = HighFidelityPdfToDocx()
        opened = pymupdf.open(str(path))
        structure = converter.analyze_page(opened[0], 0)
        opened.close()
        assert structure.column_count == 1


class TestConversion:
    def test_creates_docx_with_paragraphs(self, simple_pdf_path, tmp_path):
        output = tmp_path / "out.docx"
        pdf_to_docx_hq(str(simple_pdf_path), str(output))
        document = Document(str(output))
        texts = [paragraph.text for paragraph in document.paragraphs]
        assert any("Chapter One" in text for text in texts)

    def test_preserves_font_size_and_weight(self, simple_pdf_path, tmp_path):
        output = tmp_path / "styled.docx"
        pdf_to_docx_hq(str(simple_pdf_path), str(output))
        document = Document(str(output))

        heading = next(
            paragraph for paragraph in document.paragraphs
            if "Chapter One" in paragraph.text
        )
        run = heading.runs[0]
        assert run.bold is True
        assert abs(run.font.size.pt - 20.0) < 0.1

    def test_page_size_matches_pdf(self, simple_pdf_path, tmp_path):
        output = tmp_path / "sized.docx"
        pdf_to_docx_hq(str(simple_pdf_path), str(output))
        section = Document(str(output)).sections[0]
        assert abs(section.page_width.pt - PAGE_WIDTH) < 1.0
        assert abs(section.page_height.pt - PAGE_HEIGHT) < 1.0

    def test_multi_page_creates_sections(self, multi_page_pdf_path, tmp_path):
        output = tmp_path / "multi.docx"
        pdf_to_docx_hq(str(multi_page_pdf_path), str(output))
        document = Document(str(output))
        assert len(document.sections) == 3
        assert any("Page 1 content" in paragraph.text for paragraph in document.paragraphs)
        assert any("Page 3 content" in paragraph.text for paragraph in document.paragraphs)

    def test_header_and_footer_rendered(self, multi_page_pdf_path, tmp_path):
        output = tmp_path / "hf.docx"
        pdf_to_docx_hq(str(multi_page_pdf_path), str(output))
        document = Document(str(output))
        section = document.sections[0]
        assert "Running head 1" in section.header.paragraphs[0].text
        assert "12" in section.footer.paragraphs[0].text

    def test_table_written_to_docx(self, rich_pdf_path, tmp_path):
        output = tmp_path / "table.docx"
        pdf_to_docx_hq(str(rich_pdf_path), str(output))
        document = Document(str(output))
        assert len(document.tables) == 1
        assert len(document.tables[0].rows) == 2
        assert len(document.tables[0].columns) == 2
        assert document.tables[0].cell(0, 0).text == "Cell A1"

    def test_hyperlink_relationship_created(self, rich_pdf_path, tmp_path):
        output = tmp_path / "link.docx"
        pdf_to_docx_hq(str(rich_pdf_path), str(output))
        document = Document(str(output))

        uris = [
            rel.target_ref
            for rel in document.part.rels.values()
            if rel.reltype.endswith("/hyperlink")
        ]
        assert "https://example.com/docome" in uris

    def test_hyperlink_element_present_in_body(self, rich_pdf_path, tmp_path):
        output = tmp_path / "link2.docx"
        pdf_to_docx_hq(str(rich_pdf_path), str(output))
        document = Document(str(output))
        assert document.element.body.findall(f".//{qn('w:hyperlink')}")

    def test_image_written_to_docx(self, rich_pdf_path, tmp_path):
        output = tmp_path / "image.docx"
        pdf_to_docx_hq(str(rich_pdf_path), str(output))
        document = Document(str(output))
        assert any("media" in part.partname for part in document.part.package.parts)

    def test_list_paragraphs_use_list_styles(self, rich_pdf_path, tmp_path):
        output = tmp_path / "lists.docx"
        pdf_to_docx_hq(str(rich_pdf_path), str(output))
        document = Document(str(output))
        styles = [paragraph.style.name for paragraph in document.paragraphs]
        assert "List Bullet" in styles
        assert "List Number" in styles

    def test_multi_column_section_sets_column_count(self, rich_pdf_path, tmp_path):
        output = tmp_path / "cols.docx"
        pdf_to_docx_hq(str(rich_pdf_path), str(output))
        section = Document(str(output)).sections[0]
        columns = section._sectPr.find(qn("w:cols"))
        assert columns.get(qn("w:num")) == "2"

    def test_accepts_bytes_input(self, simple_pdf_path, tmp_path):
        payload = simple_pdf_path.read_bytes()
        output = tmp_path / "bytes.docx"
        pdf_to_docx_hq(payload, str(output))
        assert output.exists()

    def test_empty_document_produces_valid_docx(self, tmp_path):
        source = pymupdf.open()
        source.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        path = tmp_path / "blank.pdf"
        source.save(str(path))
        source.close()

        output = tmp_path / "blank.docx"
        pdf_to_docx_hq(str(path), str(output))
        assert Document(str(output)) is not None


class TestDocxBuilder:
    def test_alignment_uses_page_metrics(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        document = Document()
        builder = DocxBuilder(document)
        structure = PageStructure(page_num=0, width=PAGE_WIDTH, height=PAGE_HEIGHT)
        block = make_block("Centered", x0=250, x1=345)
        builder.add_paragraph(block, structure)
        assert document.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER

    def test_right_aligned_block_detected(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        document = Document()
        builder = DocxBuilder(document)
        structure = PageStructure(
            page_num=0, width=PAGE_WIDTH, height=PAGE_HEIGHT,
            margin_left=72.0, margin_right=72.0,
        )
        block = make_block("Right edge", x0=380, x1=523, y0=400, y1=412)
        builder.add_paragraph(block, structure)
        assert document.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT

    def test_add_table_ignores_empty(self):
        document = Document()
        builder = DocxBuilder(document)
        builder.add_table(Table(x0=0, y0=0, x1=0, y1=0, rows=0, cols=0, cells=[]))
        assert document.tables == []
