import io
import sys
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from packages.conversion_engine.docx_pdf import docx_to_html, docx_to_pdf


def make_docx_bytes(build) -> bytes:
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def render(build, tmp_path):
    source = tmp_path / "input.docx"
    target = tmp_path / "output.pdf"
    source.write_bytes(make_docx_bytes(build))
    docx_to_pdf(str(source), str(target))
    return pymupdf.open(str(target))


def all_text(document) -> str:
    return "".join(page.get_text() for page in document)


def spans_of(document, needle):
    for page in document:
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line["spans"]:
                    if needle in span["text"]:
                        yield span


class TestDocxToHtml:
    def test_paragraphs_in_order(self):
        def build(document):
            document.add_paragraph("First")
            document.add_paragraph("Second")

        markup = docx_to_html(make_docx_bytes(build))
        assert markup.index("First") < markup.index("Second")

    def test_tables_keep_document_order(self):
        def build(document):
            document.add_paragraph("Before the table")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "Alpha"
            table.rows[0].cells[1].text = "Beta"
            document.add_paragraph("After the table")

        markup = docx_to_html(make_docx_bytes(build))
        assert markup.index("Before the table") < markup.index("Alpha")
        assert markup.index("Alpha") < markup.index("After the table")
        assert "<table" in markup

    def test_run_formatting_becomes_inline_style(self):
        def build(document):
            run = document.add_paragraph().add_run("Formatted")
            run.bold = True
            run.italic = True
            run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

        markup = docx_to_html(make_docx_bytes(build))
        assert "font-weight:bold" in markup
        assert "font-style:italic" in markup
        assert "color:#ff0000" in markup

    def test_run_size_becomes_inline_style(self):
        def build(document):
            document.add_paragraph().add_run("Sized").font.size = Pt(18)

        assert "font-size:18pt" in docx_to_html(make_docx_bytes(build))

    def test_alignment_becomes_text_align(self):
        def build(document):
            paragraph = document.add_paragraph("Centered")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        assert "text-align:center" in docx_to_html(make_docx_bytes(build))

    def test_special_characters_are_escaped(self):
        def build(document):
            document.add_paragraph("5 < 6 & 7 > 2")

        markup = docx_to_html(make_docx_bytes(build))
        assert "&lt;" in markup and "&amp;" in markup
        assert "5 < 6" not in markup

    def test_accepts_bytes(self):
        def build(document):
            document.add_paragraph("From bytes")

        assert "From bytes" in docx_to_html(make_docx_bytes(build))

    def test_embedded_image_becomes_data_uri(self, tmp_path):
        image = tmp_path / "dot.png"
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20))
        pixmap.set_rect(pixmap.irect, (255, 0, 0))
        pixmap.save(str(image))

        document = Document()
        document.add_picture(str(image))

        assert "data:image/png;base64," in docx_to_html(document)


class TestDocxToPdf:
    def test_produces_readable_pdf(self, tmp_path):
        document = render(lambda d: d.add_paragraph("Readable output"), tmp_path)
        try:
            assert document.page_count == 1
            assert "Readable output" in all_text(document)
        finally:
            document.close()

    def test_uses_document_page_size_not_hardcoded_a4(self, tmp_path):
        def build(document):
            section = document.sections[0]
            section.page_width = Pt(1000)
            section.page_height = Pt(500)
            document.add_paragraph("Wide page")

        document = render(build, tmp_path)
        try:
            assert round(document[0].rect.width) == 1000
            assert round(document[0].rect.height) == 500
        finally:
            document.close()

    def test_paginates_long_documents(self, tmp_path):
        def build(document):
            for index in range(120):
                document.add_paragraph(f"Filler paragraph number {index}")

        document = render(build, tmp_path)
        try:
            assert document.page_count > 1
            assert "Filler paragraph number 119" in all_text(document)
        finally:
            document.close()

    def test_table_content_survives(self, tmp_path):
        def build(document):
            table = document.add_table(rows=2, cols=2)
            table.rows[0].cells[0].text = "Region"
            table.rows[0].cells[1].text = "Revenue"
            table.rows[1].cells[0].text = "North"
            table.rows[1].cells[1].text = "1200"

        document = render(build, tmp_path)
        try:
            text = all_text(document)
            for value in ["Region", "Revenue", "North", "1200"]:
                assert value in text
        finally:
            document.close()

    def test_bold_and_italic_faces_are_preserved(self, tmp_path):
        def build(document):
            paragraph = document.add_paragraph()
            paragraph.add_run("BoldText").bold = True
            paragraph.add_run("ItalicText").italic = True

        document = render(build, tmp_path)
        try:
            fonts = {span["font"] for span in spans_of(document, "Text")}
            assert any("Bold" in font for font in fonts)
            assert any("Italic" in font for font in fonts)
        finally:
            document.close()

    def test_run_color_is_preserved(self, tmp_path):
        def build(document):
            document.add_paragraph().add_run("Crimson").font.color.rgb = RGBColor(0xC0, 0, 0)

        document = render(build, tmp_path)
        try:
            colors = {span["color"] for span in spans_of(document, "Crimson")}
            assert 0xC00000 in colors
        finally:
            document.close()

    def test_font_size_is_preserved(self, tmp_path):
        def build(document):
            document.add_paragraph().add_run("Large").font.size = Pt(24)

        document = render(build, tmp_path)
        try:
            sizes = [span["size"] for span in spans_of(document, "Large")]
            assert sizes and round(sizes[0]) == 24
        finally:
            document.close()

    def test_embedded_image_renders(self, tmp_path):
        image = tmp_path / "dot.png"
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 40))
        pixmap.set_rect(pixmap.irect, (0, 0, 255))
        pixmap.save(str(image))

        def build(document):
            document.add_paragraph("Figure below:")
            document.add_picture(str(image))

        document = render(build, tmp_path)
        try:
            assert document[0].get_images(full=True)
            assert "Figure below:" in document[0].get_text()
        finally:
            document.close()

    def test_empty_document_still_writes_a_page(self, tmp_path):
        document = render(lambda d: None, tmp_path)
        try:
            assert document.page_count >= 1
        finally:
            document.close()

    def test_impossible_margins_raise_instead_of_hanging(self, tmp_path):
        def build(document):
            section = document.sections[0]
            section.page_width = Pt(200)
            section.page_height = Pt(200)
            section.left_margin = Pt(150)
            section.right_margin = Pt(150)
            document.add_paragraph("Does not fit")

        with pytest.raises(ValueError, match="no printable area"):
            render(build, tmp_path)

    def test_accepts_bytes_input(self, tmp_path):
        target = tmp_path / "bytes.pdf"
        docx_to_pdf(make_docx_bytes(lambda d: d.add_paragraph("Bytes in")), str(target))
        document = pymupdf.open(str(target))
        try:
            assert "Bytes in" in all_text(document)
        finally:
            document.close()


class TestEngineIntegration:
    def test_engine_docx_to_pdf_uses_renderer(self, tmp_path):
        from packages.conversion_engine import ConversionEngine

        target = tmp_path / "engine.pdf"
        ConversionEngine().docx_to_pdf(
            make_docx_bytes(lambda d: d.add_paragraph("Through the engine")), str(target)
        )
        document = pymupdf.open(str(target))
        try:
            assert "Through the engine" in all_text(document)
        finally:
            document.close()
