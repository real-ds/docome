import io
from pathlib import Path
import pytest
import pymupdf
from docx import Document

import sys
sys.path.insert(0, str(Path(__file__).parent.parent.parent))

from packages.conversion_engine import ConversionEngine


@pytest.fixture
def conversion_engine():
    return ConversionEngine()


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 72), "Test Page 1", fontsize=24)
    page.insert_text((72, 120), "This is test content.", fontsize=12)
    
    output = io.BytesIO()
    doc.save(output)
    doc.close()
    return output.getvalue()


@pytest.fixture
def sample_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Test Document", 0)
    doc.add_paragraph("This is a test paragraph.")
    doc.add_paragraph("Another paragraph.")
    
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()


class TestPdfToDocx:
    def test_pdf_to_docx(self, conversion_engine, sample_pdf_bytes, tmp_path):
        input_file = tmp_path / "input.pdf"
        output_file = tmp_path / "output.docx"
        
        with open(input_file, "wb") as f:
            f.write(sample_pdf_bytes)
        
        conversion_engine.pdf_to_docx(str(input_file), str(output_file))
        
        assert output_file.exists()
        
        doc = Document(str(output_file))
        assert len(doc.paragraphs) > 0


class TestPdfToText:
    def test_pdf_to_text(self, conversion_engine, sample_pdf_bytes, tmp_path):
        input_file = tmp_path / "input.pdf"
        
        with open(input_file, "wb") as f:
            f.write(sample_pdf_bytes)
        
        text = conversion_engine.pdf_to_text(str(input_file))
        
        assert "Test Page 1" in text
        assert "test content" in text


class TestPdfToMarkdown:
    def test_pdf_to_markdown(self, conversion_engine, sample_pdf_bytes, tmp_path):
        input_file = tmp_path / "input.pdf"
        
        with open(input_file, "wb") as f:
            f.write(sample_pdf_bytes)
        
        md = conversion_engine.pdf_to_markdown(str(input_file))
        
        assert "Test Page 1" in md


class TestDocxToPdf:
    def test_docx_to_pdf(self, conversion_engine, sample_docx_bytes, tmp_path):
        input_file = tmp_path / "input.docx"
        output_file = tmp_path / "output.pdf"
        
        with open(input_file, "wb") as f:
            f.write(sample_docx_bytes)
        
        conversion_engine.docx_to_pdf(str(input_file), str(output_file))
        
        assert output_file.exists()
        
        doc = pymupdf.open(str(output_file))
        assert len(doc) >= 1
        doc.close()


class TestPdfToXlsx:
    def test_pdf_to_xlsx(self, conversion_engine, sample_pdf_bytes, tmp_path):
        input_file = tmp_path / "input.pdf"
        output_file = tmp_path / "output.xlsx"
        
        with open(input_file, "wb") as f:
            f.write(sample_pdf_bytes)
        
        conversion_engine.pdf_to_xlsx(str(input_file), str(output_file))
        
        assert output_file.exists()
        
        import openpyxl
        wb = openpyxl.load_workbook(str(output_file))
        assert wb.active.max_row >= 1
        wb.close()
