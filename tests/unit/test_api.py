import io
import json
import sys
from pathlib import Path

import pymupdf
import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "packages"))
sys.path.insert(0, str(ROOT / "apps" / "api"))

from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def make_pdf_bytes(text="Editable content", pages=1):
    document = pymupdf.open()
    for _ in range(pages):
        page = document.new_page(width=595, height=842)
        page.insert_text((72, 100), text, fontname="helv", fontsize=12)
    buffer = io.BytesIO()
    document.save(buffer)
    document.close()
    return buffer.getvalue()


def plan_bytes(operations):
    return json.dumps({"operations": operations}).encode("utf-8")


def post_plan(pdf_bytes, operations, filename="input.pdf"):
    return client.post(
        "/api/v1/edit/apply",
        files={"file": (filename, pdf_bytes, "application/pdf")},
        data={"plan": plan_bytes(operations).decode("utf-8")},
    )


class TestHealth:
    def test_health(self):
        response = client.get("/health")
        assert response.status_code == 200
        assert response.json() == {"status": "ok"}


class TestEditActions:
    def test_lists_actions(self):
        response = client.get("/api/v1/edit/actions")
        assert response.status_code == 200
        actions = response.json()["actions"]
        assert "add_text" in actions
        assert "undo" in actions

    def test_actions_match_cli_engine(self):
        from pdf_engine import EditPlanExecutor

        response = client.get("/api/v1/edit/actions")
        assert response.json()["actions"] == EditPlanExecutor().available_actions()


class TestEditApply:
    def test_add_text_returns_pdf(self):
        response = post_plan(make_pdf_bytes(), [
            {"action": "add_text", "page": 1, "x": 100, "y": 300, "text": "Injected", "fontsize": 14},
        ])

        assert response.status_code == 200
        assert response.headers["content-type"] == "application/pdf"
        assert response.headers["x-docome-operations"] == "1"
        assert "attachment" in response.headers["content-disposition"]

        document = pymupdf.open(stream=response.content, filetype="pdf")
        assert "Injected" in document[0].get_text()
        document.close()

    def test_output_pdf_is_readable_after_response(self):
        response = post_plan(make_pdf_bytes(), [
            {"action": "add_text", "page": 1, "x": 100, "y": 300, "text": "Persisted"},
        ])
        document = pymupdf.open(stream=response.content, filetype="pdf")
        assert document.page_count == 1
        document.close()

    def test_multiple_operations(self):
        response = post_plan(make_pdf_bytes(), [
            {"action": "add_text", "page": 1, "x": 100, "y": 300, "text": "One"},
            {"action": "add_text", "page": 1, "x": 100, "y": 320, "text": "Two"},
        ])
        assert response.headers["x-docome-operations"] == "2"
        document = pymupdf.open(stream=response.content, filetype="pdf")
        text = document[0].get_text()
        document.close()
        assert "One" in text and "Two" in text

    def test_undo_action(self):
        response = post_plan(make_pdf_bytes(), [
            {"action": "add_text", "page": 1, "x": 100, "y": 300, "text": "Gone"},
            {"action": "undo"},
        ])
        document = pymupdf.open(stream=response.content, filetype="pdf")
        text = document[0].get_text()
        document.close()
        assert "Gone" not in text

    def test_annotation_operation(self):
        response = post_plan(make_pdf_bytes(), [
            {"action": "add_annotation", "page": 1, "annotation": "HIGHLIGHT",
             "x0": 70, "y0": 90, "x1": 220, "y1": 106},
        ])
        assert response.status_code == 200
        document = pymupdf.open(stream=response.content, filetype="pdf")
        assert len(list(document[0].annots() or [])) == 1
        document.close()

    def test_empty_plan_rejected(self):
        response = post_plan(make_pdf_bytes(), [])
        assert response.status_code == 422

    def test_unknown_action_rejected(self):
        response = post_plan(make_pdf_bytes(), [{"action": "fly", "page": 1}])
        assert response.status_code == 400
        assert "unknown action" in response.json()["detail"]

    def test_invalid_page_rejected(self):
        response = post_plan(make_pdf_bytes(), [
            {"action": "add_text", "page": 99, "x": 1, "y": 1, "text": "t"},
        ])
        assert response.status_code == 400
        assert "Invalid page" in response.json()["detail"]

    def test_missing_plan_field_rejected(self):
        response = client.post(
            "/api/v1/edit/apply",
            files={"file": ("input.pdf", make_pdf_bytes(), "application/pdf")},
        )
        assert response.status_code == 422

    def test_malformed_plan_json_rejected(self):
        response = client.post(
            "/api/v1/edit/apply",
            files={"file": ("input.pdf", make_pdf_bytes(), "application/pdf")},
            data={"plan": "{not json"},
        )
        assert response.status_code == 422

    def test_operation_missing_action_rejected(self):
        response = post_plan(make_pdf_bytes(), [{"page": 1, "x": 1, "y": 1}])
        assert response.status_code == 422

    def test_negative_page_rejected_by_schema(self):
        response = post_plan(make_pdf_bytes(), [
            {"action": "add_text", "page": -1, "x": 1, "y": 1, "text": "t"},
        ])
        assert response.status_code == 422

    def test_unknown_fields_are_dropped(self):
        response = post_plan(make_pdf_bytes(), [
            {"action": "add_text", "page": 1, "x": 100, "y": 300, "text": "Clean",
             "totally_unknown": "ignored"},
        ])
        assert response.status_code == 200

    def test_empty_upload_rejected(self):
        response = post_plan(b"", [
            {"action": "add_text", "page": 1, "x": 1, "y": 1, "text": "t"},
        ])
        assert response.status_code == 400

    def test_path_traversal_filename_is_sanitized(self):
        response = post_plan(make_pdf_bytes(), [
            {"action": "add_text", "page": 1, "x": 100, "y": 300, "text": "Safe"},
        ], filename="../../etc/passwd.pdf")
        assert response.status_code == 200
        assert "/" not in response.headers["content-disposition"].split("filename=")[1]

    def test_output_can_reuse_the_upload_filename(self):
        response = post_plan(make_pdf_bytes(), [
            {"action": "add_text", "page": 1, "x": 100, "y": 300, "text": "SameName"},
        ], filename="input.pdf")
        assert response.status_code == 200
        document = pymupdf.open(stream=response.content, filetype="pdf")
        assert "SameName" in document[0].get_text()
        document.close()

    def test_docx_route_can_reuse_upload_filename(self):
        response = client.post(
            "/api/v1/convert/pdf2docx-hq",
            files={"file": ("report.pdf", make_pdf_bytes("Body"), "application/pdf")},
        )
        assert response.status_code == 200


class TestHighFidelityConversion:
    def test_converts_to_docx(self):
        response = client.post(
            "/api/v1/convert/pdf2docx-hq",
            files={"file": ("report.pdf", make_pdf_bytes("Reconstruction target"), "application/pdf")},
        )
        assert response.status_code == 200
        assert "wordprocessingml" in response.headers["content-type"]

        from docx import Document

        document = Document(io.BytesIO(response.content))
        texts = [paragraph.text for paragraph in document.paragraphs]
        assert any("Reconstruction target" in text for text in texts)

    def test_empty_upload_rejected(self):
        response = client.post(
            "/api/v1/convert/pdf2docx-hq",
            files={"file": ("empty.pdf", b"", "application/pdf")},
        )
        assert response.status_code == 400

    def test_non_pdf_upload_rejected(self):
        response = client.post(
            "/api/v1/convert/pdf2docx-hq",
            files={"file": ("bad.pdf", b"not a pdf at all", "application/pdf")},
        )
        assert response.status_code == 422


class TestFidelityEndpoint:
    def _docx_for(self, pdf_bytes):
        conversion = client.post(
            "/api/v1/convert/pdf2docx-hq",
            files={"file": ("report.pdf", pdf_bytes, "application/pdf")},
        )
        assert conversion.status_code == 200
        return conversion.content

    def test_scores_conversion(self):
        pdf_bytes = make_pdf_bytes("Fidelity subject line")
        response = client.post(
            "/api/v1/convert/fidelity",
            files={
                "file": ("report.pdf", pdf_bytes, "application/pdf"),
                "converted": ("report.docx", self._docx_for(pdf_bytes),
                              "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            },
        )
        assert response.status_code == 200
        payload = response.json()
        assert payload["text_retention"] == 1.0
        assert payload["page_count_preserved"] is True
        assert payload["failed_metrics"] == []

    def test_reports_failed_metrics(self):
        from docx import Document

        stripped = Document()
        stripped.add_paragraph("completely unrelated")
        buffer = io.BytesIO()
        stripped.save(buffer)

        response = client.post(
            "/api/v1/convert/fidelity",
            files={
                "file": ("report.pdf", make_pdf_bytes(), "application/pdf"),
                "converted": ("report.docx", buffer.getvalue(),
                              "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            },
        )
        assert response.status_code == 200
        payload = response.json()
        assert payload["passed"] is False
        assert "text_retention" in payload["failed_metrics"]

    def test_missing_converted_file_rejected(self):
        response = client.post(
            "/api/v1/convert/fidelity",
            files={"file": ("report.pdf", make_pdf_bytes(), "application/pdf")},
        )
        assert response.status_code == 422

    def test_invalid_docx_rejected(self):
        response = client.post(
            "/api/v1/convert/fidelity",
            files={
                "file": ("report.pdf", make_pdf_bytes(), "application/pdf"),
                "converted": ("report.docx", b"not a docx", "application/octet-stream"),
            },
        )
        assert response.status_code == 422


class TestPdfOperations:
    def test_merge_two_pdfs(self):
        response = client.post(
            "/api/v1/pdf/merge",
            files=[
                ("files", ("a.pdf", make_pdf_bytes("First"), "application/pdf")),
                ("files", ("b.pdf", make_pdf_bytes("Second"), "application/pdf")),
            ],
        )
        assert response.status_code == 200
        document = pymupdf.open(stream=response.content, filetype="pdf")
        assert document.page_count == 2
        document.close()

    def test_merge_requires_two_files(self):
        response = client.post(
            "/api/v1/pdf/merge",
            files=[("files", ("a.pdf", make_pdf_bytes("Only"), "application/pdf"))],
        )
        assert response.status_code == 400

    def test_split_single_page_returns_pdf(self):
        response = client.post(
            "/api/v1/pdf/split",
            files={"file": ("doc.pdf", make_pdf_bytes("Split me", pages=2), "application/pdf")},
            data={"pages_per_split": "2"},
        )
        assert response.status_code == 200
        document = pymupdf.open(stream=response.content, filetype="pdf")
        assert document.page_count == 2
        document.close()

    def test_split_many_pages_returns_zip(self):
        response = client.post(
            "/api/v1/pdf/split",
            files={"file": ("doc.pdf", make_pdf_bytes("Split me", pages=3), "application/pdf")},
            data={"pages_per_split": "1"},
        )
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/zip"

        import zipfile

        archive = zipfile.ZipFile(io.BytesIO(response.content))
        assert len(archive.namelist()) == 3

    def test_split_rejects_zero(self):
        response = client.post(
            "/api/v1/pdf/split",
            files={"file": ("doc.pdf", make_pdf_bytes(), "application/pdf")},
            data={"pages_per_split": "0"},
        )
        assert response.status_code == 400

    def test_extract_pages(self):
        response = client.post(
            "/api/v1/pdf/extract",
            files={"file": ("doc.pdf", make_pdf_bytes("Extract", pages=3), "application/pdf")},
            data={"pages": "1,3"},
        )
        assert response.status_code == 200
        document = pymupdf.open(stream=response.content, filetype="pdf")
        assert document.page_count == 2
        document.close()

    def test_extract_requires_selection(self):
        response = client.post(
            "/api/v1/pdf/extract",
            files={"file": ("doc.pdf", make_pdf_bytes(), "application/pdf")},
        )
        assert response.status_code == 400

    def test_extract_rejects_bad_pages(self):
        response = client.post(
            "/api/v1/pdf/extract",
            files={"file": ("doc.pdf", make_pdf_bytes(), "application/pdf")},
            data={"pages": "one,two"},
        )
        assert response.status_code == 400

    def test_extract_rejects_bad_range(self):
        response = client.post(
            "/api/v1/pdf/extract",
            files={"file": ("doc.pdf", make_pdf_bytes(), "application/pdf")},
            data={"start": "3", "end": "1"},
        )
        assert response.status_code == 400

    def test_rotate(self):
        response = client.post(
            "/api/v1/pdf/rotate",
            files={"file": ("doc.pdf", make_pdf_bytes(), "application/pdf")},
            data={"degrees": "90"},
        )
        assert response.status_code == 200
        document = pymupdf.open(stream=response.content, filetype="pdf")
        assert document[0].rotation == 90
        document.close()

    def test_rotate_rejects_bad_degrees(self):
        response = client.post(
            "/api/v1/pdf/rotate",
            files={"file": ("doc.pdf", make_pdf_bytes(), "application/pdf")},
            data={"degrees": "45"},
        )
        assert response.status_code == 400

    def test_compress(self):
        response = client.post(
            "/api/v1/pdf/compress",
            files={"file": ("doc.pdf", make_pdf_bytes("Compress me", pages=3), "application/pdf")},
            data={"level": "recommended"},
        )
        assert response.status_code == 200
        document = pymupdf.open(stream=response.content, filetype="pdf")
        assert document.page_count == 3
        document.close()

    def test_compress_rejects_bad_level(self):
        response = client.post(
            "/api/v1/pdf/compress",
            files={"file": ("doc.pdf", make_pdf_bytes(), "application/pdf")},
            data={"level": "turbo"},
        )
        assert response.status_code == 400

    def test_get_metadata(self):
        response = client.post(
            "/api/v1/pdf/metadata/read",
            files={"file": ("doc.pdf", make_pdf_bytes(), "application/pdf")},
        )
        assert response.status_code == 200
        assert set(response.json()) >= {"title", "author", "producer"}

    def test_update_metadata(self):
        response = client.post(
            "/api/v1/pdf/metadata",
            files={"file": ("doc.pdf", make_pdf_bytes(), "application/pdf")},
            data={"metadata": json.dumps({"title": "Quarterly Report", "author": "Docome"})},
        )
        assert response.status_code == 200
        document = pymupdf.open(stream=response.content, filetype="pdf")
        assert document.metadata["title"] == "Quarterly Report"
        document.close()

    def test_update_metadata_rejects_bad_json(self):
        response = client.post(
            "/api/v1/pdf/metadata",
            files={"file": ("doc.pdf", make_pdf_bytes(), "application/pdf")},
            data={"metadata": "{oops"},
        )
        assert response.status_code == 422

    def test_page_count(self):
        response = client.post(
            "/api/v1/pdf/pages/count",
            files={"file": ("doc.pdf", make_pdf_bytes(pages=4), "application/pdf")},
        )
        assert response.status_code == 200
        assert response.json() == {"page_count": 4}

    def test_empty_upload_rejected(self):
        response = client.post(
            "/api/v1/pdf/pages/count",
            files={"file": ("doc.pdf", b"", "application/pdf")},
        )
        assert response.status_code == 400


class TestBasicConversionRoutes:
    def test_pdf2txt(self):
        response = client.post(
            "/api/v1/convert/pdf2txt",
            files={"file": ("doc.pdf", make_pdf_bytes("Plain text target"), "application/pdf")},
        )
        assert response.status_code == 200
        assert "Plain text target" in response.json()["text"]

    def test_pdf2md(self):
        response = client.post(
            "/api/v1/convert/pdf2md",
            files={"file": ("doc.pdf", make_pdf_bytes("Markdown target"), "application/pdf")},
        )
        assert response.status_code == 200
        assert "Markdown target" in response.json()["markdown"]

    def test_pdf2docx(self):
        response = client.post(
            "/api/v1/convert/pdf2docx",
            files={"file": ("doc.pdf", make_pdf_bytes("Basic conversion"), "application/pdf")},
        )
        assert response.status_code == 200
        assert "wordprocessingml" in response.headers["content-type"]

    def test_empty_upload_rejected(self):
        response = client.post(
            "/api/v1/convert/pdf2txt",
            files={"file": ("doc.pdf", b"", "application/pdf")},
        )
        assert response.status_code == 400


class TestWorkspaceHelper:
    def test_safe_name_strips_traversal(self):
        from app.core.workspace import safe_name

        assert "/" not in safe_name("../../etc/passwd")
        assert safe_name("") == "upload"
        assert safe_name(None) == "upload"

    def test_workspace_cleans_up(self):
        from app.core.workspace import workspace

        with workspace() as directory:
            created = directory / "file.txt"
            created.write_text("data")
            assert created.exists()
            captured = directory

        assert not captured.exists()

    def test_uploaded_file_removed(self):
        from app.core.workspace import uploaded_file, workspace

        with workspace() as directory:
            with uploaded_file(directory, "in.pdf", b"bytes") as path:
                assert path.read_bytes() == b"bytes"
                captured = path
            assert not captured.exists()
